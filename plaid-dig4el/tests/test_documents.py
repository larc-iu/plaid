"""Offline checks of the reference-document helpers: text extraction from the three
file kinds and chunking."""

import io

from docx import Document
from pypdf import PdfWriter

from plaid_dig4el import documents


def test_extract_text_from_docx_and_txt():
    d = Document()
    d.add_paragraph("Negation uses the particle 'aita.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "word", "gloss"
    buf = io.BytesIO(); d.save(buf)
    text = documents.extract_text("g.docx", buf.getvalue())
    assert "particle 'aita" in text and "word | gloss" in text
    assert documents.extract_text("notes.md", "# Notes\nplain".encode()) == "# Notes\nplain"


def test_extract_text_from_pdf_does_not_crash_on_blank_pages():
    w = PdfWriter(); w.add_blank_page(width=200, height=200)
    buf = io.BytesIO(); w.write(buf)
    assert documents.extract_text("blank.pdf", buf.getvalue()).strip() == ""


def test_extract_text_rejects_unknown_kinds():
    try:
        documents.extract_text("x.xlsx", b"...")
    except ValueError as e:
        assert "PDF, DOCX or plain text" in str(e)
    else:
        raise AssertionError("expected a ValueError")


def test_chunks_overlap_and_cover_the_text():
    text = "\n\n".join(f"Paragraph {i}. " + ("word " * 120).strip() + "." for i in range(12))
    chunks = documents.chunk_text(text, size=800, overlap=100)
    assert len(chunks) > 3
    assert all(len(c) <= 800 for c in chunks)
    assert chunks[0].startswith("Paragraph 0.") and "Paragraph 11." in chunks[-1]
    # consecutive chunks share text (the overlap)
    assert any(chunks[i][-40:] in chunks[i + 1] or chunks[i + 1][:40] in chunks[i] for i in range(len(chunks) - 1))
    assert documents.chunk_text("") == []
