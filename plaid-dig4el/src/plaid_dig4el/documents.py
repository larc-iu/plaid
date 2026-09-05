"""Reference documents: files about the language that generation can draw on.

dig4el uploaded PDFs to an OpenAI vector store and asked gpt-4.1 with file search to
compile an answer (``file_search_request``). Here the file stays on disk, its text is
extracted and cut into chunks embedded on the configured endpoint, the chunks nearest
the query are retrieved, and the same prompt (``generation.DOCUMENTS_PROMPT``) is
answered over them. The outcome has the same shape: an answer text and the names of
the documents it drew on.
"""

from __future__ import annotations

import io
from pathlib import Path

import numpy as np

from . import db, jobs
from .config import settings
from .llm import LLM, llm

CHUNK_CHARS = 1500
CHUNK_OVERLAP = 200
RETRIEVE_K = 12
TEXT_SUFFIXES = {".txt", ".md", ".csv", ".json", ".tex", ".html", ".htm"}


def documents_dir(language_id: str) -> Path:
    d = settings().data_dir / "documents" / language_id
    d.mkdir(parents=True, exist_ok=True)
    return d


def stored_path(doc: db.ReferenceDocument) -> Path:
    return documents_dir(doc.language_id) / f"{doc.id}-{doc.filename}"


def extract_text(filename: str, raw: bytes) -> str:
    """Plain text from a PDF, a Word document, or a text file."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == ".docx":
        from docx import Document

        d = Document(io.BytesIO(raw))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if suffix in TEXT_SUFFIXES or not suffix:
        return raw.decode("utf-8", errors="replace")
    raise ValueError(f"Cannot read {suffix} files. Use PDF, DOCX or plain text.")


def chunk_text(text: str, size: int = CHUNK_CHARS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Overlapping windows over the text, cut at paragraph or sentence ends when possible."""
    text = text.replace("\r", "")
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        if end < len(text):
            cut = max(text.rfind("\n\n", start + size // 2, end), text.rfind(". ", start + size // 2, end))
            if cut > start:
                end = cut + 1
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def add_document(language_id: str, filename: str, raw: bytes, title: str, description: str,
                 uploaded_by: str) -> db.ReferenceDocument:
    """Store the file and its text; indexing happens in the ``index_documents`` job."""
    doc = db.ReferenceDocument(language_id=language_id, filename=filename, title=title or filename,
                               description=description, size=len(raw), uploaded_by=uploaded_by)
    try:
        doc.text = extract_text(filename, raw)
    except Exception as e:  # unreadable file: kept, marked
        doc.status = "failed"
        doc.error = str(e)[:500]
    with db.session() as s:
        s.add(doc)
        s.commit()
        stored_path(doc).write_bytes(raw)
        return doc


def remove_document(doc_id: str) -> None:
    with db.session() as s:
        doc = s.get(db.ReferenceDocument, doc_id)
        if doc is None:
            return
        path = stored_path(doc)
        s.delete(doc)
        s.commit()
    if path.exists():
        path.unlink()


@jobs.handler("index_documents")
def index_documents(job: db.Job, client) -> None:
    """Embed the chunks of every uploaded, not yet indexed document of the language."""
    language_id = job.payload["language_id"]
    L = llm()
    with db.session() as s:
        todo = [d.id for d in s.query(db.ReferenceDocument)
                .filter_by(language_id=language_id, status="uploaded").all()]
    for n, doc_id in enumerate(todo, start=1):
        with db.session() as s:
            doc = s.get(db.ReferenceDocument, doc_id)
            pieces = chunk_text(doc.text)
            try:
                vectors = L.embed(pieces) if pieces else []
            except Exception as e:
                doc.status, doc.error = "failed", str(e)[:500]
                s.commit()
                continue
            doc.chunks = [db.DocumentChunk(index=i, text=t, vector=_pack(v))
                          for i, (t, v) in enumerate(zip(pieces, vectors))]
            doc.status = "indexed"
            doc.error = ""
            s.commit()
        with db.session() as s:
            j = s.get(db.Job, job.id)
            j.progress = {"done": n, "total": len(todo), "note": ""}
            s.commit()


def _pack(v: list[float]) -> bytes:
    arr = np.asarray(v, dtype=np.float32)
    return (arr / (np.linalg.norm(arr) or 1)).tobytes()


def retrieve(language_id: str, query: str, k: int = RETRIEVE_K) -> list[tuple[db.DocumentChunk, str, float]]:
    """The chunks nearest the query across the language's indexed documents:
    (chunk, document filename, cosine score)."""
    L = llm()
    with db.session() as s:
        rows = [(c, d.filename) for d in s.query(db.ReferenceDocument)
                .filter_by(language_id=language_id, status="indexed").all() for c in d.chunks if c.vector]
    if not rows:
        return []
    q = np.asarray(L.embed([query])[0], dtype=np.float32)
    q /= np.linalg.norm(q) or 1
    matrix = np.stack([np.frombuffer(c.vector, dtype=np.float32) for c, _ in rows])
    scores = matrix @ q
    order = np.argsort(-scores)[:k]
    return [(rows[i][0], rows[i][1], float(scores[i])) for i in order]


def contribute(L: LLM, language: db.Language, query: str, model: str | None = None) -> dict | None:
    """dig4el's documents contribution: its file-search prompt answered over the
    retrieved chunks. None when the language has no indexed document."""
    from .generation import DOCUMENTS_PROMPT

    hits = retrieve(language.id, query)
    if not hits:
        return None
    passages = "\n\n".join(f"[{name}, part {c.index + 1}]\n{c.text}" for c, name, _ in hits)
    prompt = DOCUMENTS_PROMPT.format(indi_language=language.name, query=query)
    text = L.chat_text([{"role": "system", "content": prompt},
                        {"role": "user", "content": f"DOCUMENTS:\n\n{passages}"}],
                       model=model or L.model, max_tokens=6000)
    return {"text": text, "sources": sorted({name for _, name, _ in hits})}
